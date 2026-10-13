"""
PCAP Analysis Report Generator

Creates JSON and HTML reports for individual PCAP analysis runs.
"""

import json
from pathlib import Path
from datetime import datetime


def _safe_value(value):
    """Convert values into JSON-safe Python values."""

    if hasattr(value, "item"):
        return value.item()

    if isinstance(value, Path):
        return str(value)

    return value


def save_json_report(
    report_data: dict,
    output_path: Path,
) -> Path:
    """Save the complete analysis report as JSON."""

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    with open(
        output_path,
        "w",
        encoding="utf-8",
    ) as f:
        json.dump(
            report_data,
            f,
            indent=2,
            default=_safe_value,
        )

    return output_path


def _html_escape(value):
    """Escape text for safe HTML output."""

    text = str(value)

    return (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
        .replace("'", "&#39;")
    )


def _table_rows(data: dict) -> str:
    """Convert a dictionary into HTML table rows."""

    rows = []

    for key, value in data.items():

        if isinstance(value, float):
            value = f"{value:.4f}"

        rows.append(
            f"""
            <tr>
                <td>{_html_escape(key)}</td>
                <td>{_html_escape(value)}</td>
            </tr>
            """
        )

    return "\n".join(rows)


def save_html_report(
    report_data: dict,
    output_path: Path,
) -> Path:
    """Save the complete analysis report as HTML."""

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    pcap = report_data.get(
        "pcap",
        {},
    )

    extraction = report_data.get(
        "extraction",
        {},
    )

    features = report_data.get(
        "features",
        {},
    )

    model = report_data.get(
        "model",
        {},
    )

    predictions = report_data.get(
        "predictions",
        {},
    )

    ground_truth = report_data.get(
        "ground_truth"
    )

    evaluation = report_data.get(
        "evaluation"
    )

    inference = report_data.get(
        "inference"
    )

    multiclass = report_data.get(
        "multiclass"
    )

    multiclass_evaluation = report_data.get(
        "multiclass_evaluation"
    )

    generated_at = report_data.get(
        "generated_at",
        datetime.now().isoformat(),
    )

    html = f"""<!DOCTYPE html>
<html lang="en">

<head>

<meta charset="UTF-8">

<meta
    name="viewport"
    content="width=device-width, initial-scale=1.0"
>

<title>PCAP Analysis Report</title>

<style>

body {{
    font-family:
        Arial,
        Helvetica,
        sans-serif;

    margin: 0;
    padding: 0;

    background: #f5f7fa;
    color: #1f2937;
}}

.container {{
    max-width: 1100px;

    margin: 40px auto;

    background: white;

    padding: 40px;

    border-radius: 12px;
}}

h1 {{
    margin-bottom: 5px;
}}

h2 {{
    margin-top: 35px;

    border-bottom:
        2px solid #e5e7eb;

    padding-bottom: 8px;
}}

.subtitle {{
    color: #6b7280;

    margin-bottom: 30px;
}}

.summary {{
    display: grid;

    grid-template-columns:
        repeat(
            auto-fit,
            minmax(180px, 1fr)
        );

    gap: 15px;

    margin: 25px 0;
}}

.card {{
    background: #f9fafb;

    border:
        1px solid #e5e7eb;

    border-radius: 10px;

    padding: 18px;
}}

.card-title {{
    font-size: 13px;

    color: #6b7280;
}}

.card-value {{
    font-size: 25px;

    font-weight: bold;

    margin-top: 5px;
}}

table {{
    width: 100%;

    border-collapse: collapse;

    margin-top: 15px;
}}

th,
td {{
    text-align: left;

    padding: 10px;

    border-bottom:
        1px solid #e5e7eb;
}}

th {{
    background: #f3f4f6;
}}

.footer {{
    margin-top: 40px;

    padding-top: 20px;

    border-top:
        1px solid #e5e7eb;

    color: #6b7280;

    font-size: 13px;
}}

</style>

</head>

<body>

<div class="container">

<h1>
    PCAP Analysis Report
</h1>

<div class="subtitle">
    Generated:
    {_html_escape(generated_at)}
</div>


<div class="summary">


<div class="card">

<div class="card-title">
    Packets
</div>

<div class="card-value">
    {_html_escape(
        extraction.get(
            "packet_count",
            "N/A"
        )
    )}
</div>

</div>


<div class="card">

<div class="card-title">
    Flows
</div>

<div class="card-value">
    {_html_escape(
        extraction.get(
            "flow_count",
            "N/A"
        )
    )}
</div>

</div>


<div class="card">

<div class="card-title">
    ML Features
</div>

<div class="card-value">
    {_html_escape(
        features.get(
            "feature_count",
            "N/A"
        )
    )}
</div>

</div>


<div class="card">

<div class="card-title">
    Attack Predictions
</div>

<div class="card-value">
    {_html_escape(
        predictions.get(
            "attack_count",
            "N/A"
        )
    )}
</div>

</div>


<div class="card">

<div class="card-title">
    Attack Categories
</div>

<div class="card-value">
    {_html_escape(
        multiclass.get("attack_categories_detected_count", 0) if multiclass else "N/A"
    )}
</div>

</div>


<div class="card">

<div class="card-title">
    Uncertain Flows
</div>

<div class="card-value">
    {_html_escape(
        multiclass.get("uncertain_count", 0) if multiclass else "N/A"
    )}
</div>

</div>


</div>


<h2>
    1. PCAP Information
</h2>

<table>

<tr>
    <th>Property</th>
    <th>Value</th>
</tr>

{_table_rows(pcap)}

</table>


<h2>
    2. Packet and Flow Extraction
</h2>

<table>

<tr>
    <th>Property</th>
    <th>Value</th>
</tr>

{_table_rows(extraction)}

</table>


<h2>
    3. Feature Engineering
</h2>

<table>

<tr>
    <th>Property</th>
    <th>Value</th>
</tr>

{_table_rows(features)}

</table>


<h2>
    4. TabPFN Model
</h2>

<table>

<tr>
    <th>Property</th>
    <th>Value</th>
</tr>

{_table_rows(model)}

</table>
"""

    if inference is not None:
        html += f"""

<h2>
    Inference Architecture &amp; Execution
</h2>

<table>

<tr>
    <th>Property</th>
    <th>Value</th>
</tr>

{_table_rows(inference)}

</table>
"""

    html += f"""

<h2>
    5. Prediction Results
</h2>

<table>

<tr>
    <th>Property</th>
    <th>Value</th>
</tr>

{_table_rows(predictions)}

</table>
"""

    if multiclass is not None:
        category_counts = multiclass.get("category_counts", {})
        category_percentages = multiclass.get("category_percentages", {})
        total_p = predictions.get("total", 1) or 1

        cat_rows = []
        for cat, cnt in category_counts.items():
            pct = category_percentages.get(cat, round((cnt / total_p) * 100, 2))
            is_atk = cat.lower() != "normal"
            bar_color = "#ef4444" if is_atk else "#22c55e"
            cat_rows.append(f"""
            <tr>
                <td><strong>{_html_escape(cat)}</strong></td>
                <td>{_html_escape(cnt)}</td>
                <td>{pct:.2f}%</td>
                <td style="width: 35%;">
                    <div style="background:#e5e7eb;border-radius:4px;height:12px;width:100%;overflow:hidden;">
                        <div style="background:{bar_color};height:100%;width:{min(100.0, pct)}%;"></div>
                    </div>
                </td>
            </tr>
            """)
        cat_table_body = "\n".join(cat_rows)

        status_counts = multiclass.get("status_counts", {})
        conf_summary = multiclass.get("avg_confidence")
        conf_str = f"{conf_summary * 100:.2f}%" if conf_summary is not None else "N/A"
        thresh = multiclass.get("uncertainty_threshold", 0.5)

        html += f"""

<h2>
    6. Multiclass Attack Category Distribution
</h2>

<table>
<tr>
    <th>Attack Category</th>
    <th>Flow Count</th>
    <th>Percentage</th>
    <th>Distribution</th>
</tr>
{cat_table_body}
</table>

<h2>
    7. Classification Confidence &amp; Status
</h2>

<table>
<tr>
    <th>Metric</th>
    <th>Value</th>
</tr>
<tr>
    <td>Status: Classified</td>
    <td>{_html_escape(status_counts.get("classified", "N/A"))}</td>
</tr>
<tr>
    <td>Status: Uncertain (&lt; {thresh:.2f} confidence)</td>
    <td>{_html_escape(status_counts.get("uncertain", "N/A"))}</td>
</tr>
<tr>
    <td>Status: Unknown / Unclassifiable</td>
    <td>{_html_escape(status_counts.get("unknown", "N/A"))}</td>
</tr>
<tr>
    <td>Uncertain Rate</td>
    <td>{multiclass.get("uncertain_rate_percentage", 0.0):.2f}%</td>
</tr>
<tr>
    <td>Average Attack Category Confidence</td>
    <td>{conf_str}</td>
</tr>
</table>
<p style="font-size:12px;color:#6b7280;margin-top:8px;">
    <strong>Model Class Support:</strong> 9 trained categories (Normal, Backdoors, DoS, Exploits, Fuzzers, Generic, Reconnaissance, Shellcode, Worms). The official UNSW-NB15 category <em>Analysis</em> is unsupported by this model because 0 flow samples existed in the training slice (capture window 1421927376–1421928304).
</p>

<div style="background:#f0fdf4;border-left:4px solid #16a34a;padding:14px 18px;border-radius:8px;margin:25px 0;font-size:13px;color:#166534;line-height:1.5;">
    <strong>Investigative Guidance:</strong> Predictions generated by TabPFN-NIDS are machine-learning classifications based on 67 bidirectional network-flow features. Predicted attack categories indicate behavioral similarity to known threat signatures in the UNSW-NB15 benchmark, and should be treated as automated investigative triage rather than conclusive forensic proof.
</div>
"""

    if multiclass_evaluation is not None:
        mc_metrics_dict = {
            "Overall Accuracy": multiclass_evaluation.get("accuracy", "N/A"),
            "Macro Precision": multiclass_evaluation.get("macro_precision", "N/A"),
            "Macro Recall": multiclass_evaluation.get("macro_recall", "N/A"),
            "Macro F1-Score": multiclass_evaluation.get("macro_f1", "N/A"),
            "Weighted F1-Score": multiclass_evaluation.get("weighted_f1", "N/A"),
            "OvR ROC-AUC": multiclass_evaluation.get("roc_auc_ovr", "N/A"),
            "Prediction Coverage": f"{multiclass_evaluation.get('prediction_coverage', 1.0) * 100:.2f}%" if multiclass_evaluation.get('prediction_coverage') is not None else "N/A",
            "Uncertain Rate": f"{multiclass_evaluation.get('uncertain_rate', 0.0) * 100:.2f}%" if multiclass_evaluation.get('uncertain_rate') is not None else "N/A",
        }

        per_class = multiclass_evaluation.get("per_class", {})
        pc_rows = []
        for cls_name, cls_m in per_class.items():
            p_val = cls_m.get("precision", 0.0)
            r_val = cls_m.get("recall", 0.0)
            f_val = cls_m.get("f1", 0.0)
            s_val = cls_m.get("support", 0)
            pc_rows.append(f"""
            <tr>
                <td><strong>{_html_escape(cls_name)}</strong></td>
                <td>{p_val:.4f}</td>
                <td>{r_val:.4f}</td>
                <td>{f_val:.4f}</td>
                <td>{s_val}</td>
            </tr>
            """)
        pc_table_body = "\n".join(pc_rows)

        html += f"""
<h2>
    Multiclass Ground-Truth Evaluation
</h2>

<table>
<tr>
    <th>Multiclass Metric</th>
    <th>Value</th>
</tr>
{_table_rows(mc_metrics_dict)}
</table>

<h3>Per-Class Performance</h3>
<table>
<tr>
    <th>Category</th>
    <th>Precision</th>
    <th>Recall</th>
    <th>F1-Score</th>
    <th>Support</th>
</tr>
{pc_table_body}
</table>
"""

    if ground_truth is not None:

        html += f"""

<h2>
    6. Ground Truth Matching
</h2>

<table>

<tr>
    <th>Property</th>
    <th>Value</th>
</tr>

{_table_rows(ground_truth)}

</table>
"""

    if evaluation is not None:

        html += f"""

<h2>
    7. Evaluation Metrics
</h2>

<table>

<tr>
    <th>Metric</th>
    <th>Value</th>
</tr>

{_table_rows(evaluation)}

</table>
"""

    html += """

<div class="footer">

Generated by the
TabPFN-NIDS PCAP Analysis Pipeline.

</div>

</div>

</body>

</html>
"""

    with open(
        output_path,
        "w",
        encoding="utf-8",
    ) as f:
        f.write(html)

    return output_path


def build_report_data(
    pcap_path,
    extraction_summary,
    features,
    feature_names,
    predictions,
    probabilities,
    total_analysis_seconds,
    ground_truth_path=None,
    metrics=None,
    labeled_flows=None,
    inference_meta=None,
    multiclass_predictions=None,
    multiclass_probabilities=None,
    predicted_categories=None,
    category_confidences=None,
    prediction_statuses=None,
    multiclass_metrics=None,
    class_names=None,
    uncertainty_threshold=0.5,
    performance_breakdown=None,
):
    """Build structured data for JSON and HTML reports."""

    total_predictions = len(
        predictions
    )

    normal_count = int(
        (predictions == 0).sum()
    )

    attack_count = int(
        (predictions == 1).sum()
    )

    normal_percentage = (
        normal_count
        / total_predictions
        * 100
        if total_predictions
        else 0.0
    )

    attack_percentage = (
        attack_count
        / total_predictions
        * 100
        if total_predictions
        else 0.0
    )

    pcap_size_bytes = (
        pcap_path.stat().st_size
    )

    report_data = {

        "generated_at":
            datetime.now().isoformat(),

        "pcap": {

            "filename":
                pcap_path.name,

            "path":
                str(pcap_path),

            "size_bytes":
                pcap_size_bytes,

            "size_mb":
                round(
                    pcap_size_bytes
                    / (1024 * 1024),
                    4,
                ),
        },

        "extraction": {

            "packet_count":
                extraction_summary.get(
                    "total_packets",
                    0,
                ),

            "flow_count":
                extraction_summary.get(
                    "total_flows",
                    0,
                ),

            "backend":
                "scapy",

            "flow_timeout_seconds":
                120,

            "idle_timeout_seconds":
                60,
        },

        "features": {

            "flow_rows":
                len(features),

            "feature_count":
                len(feature_names),

            "feature_columns":
                list(feature_names),

            "total_columns_after_engineering":
                len(features.columns),
        },

        "model": {

            "model_type":
                "TabPFN",

            "task":
                "binary",

            "device":
                "CPU",

            "feature_count":
                len(feature_names),

            "samples_predicted":
                total_predictions,
        },

        "predictions": {

            "total":
                total_predictions,

            "normal_count":
                normal_count,

            "attack_count":
                attack_count,

            "normal_percentage":
                round(
                    normal_percentage,
                    2,
                ),

            "attack_percentage":
                round(
                    attack_percentage,
                    2,
                ),
        },

        "performance": {
            "total_analysis_seconds": round(total_analysis_seconds, 2),
            "flows_per_second": round(total_predictions / total_analysis_seconds, 2) if total_analysis_seconds > 0 else 0.0,
            "flows_processed": total_predictions,
            **(performance_breakdown or {}),
        },

        "ground_truth":
            None,

        "evaluation":
            None,

        "inference":
            None,
    }

    if inference_meta is not None:
        rows_per_chunk = inference_meta.get("rows_per_chunk", [])
        if isinstance(rows_per_chunk, list):
            rows_per_chunk_str = ", ".join(str(r) for r in rows_per_chunk)
        else:
            rows_per_chunk_str = str(rows_per_chunk)

        inf_sec = inference_meta.get("total_inference_seconds", 0.0)
        inf_rps = inference_meta.get("rows_per_second", 0.0)
        if (inf_rps == 0.0 or inf_rps < 1.0) and inf_sec > 0:
            inf_rps = round(total_predictions / inf_sec, 2)

        report_data["inference"] = {
            "Inference mode":
                inference_meta.get("inference_mode", "Single model"),
            "Number of input flows":
                inference_meta.get("num_flows", total_predictions),
            "Number of models":
                inference_meta.get("num_models", 1),
            "Number of workers":
                inference_meta.get("num_workers", 1),
            "Chunk size":
                inference_meta.get("max_rows_per_worker", 10000),
            "Number of chunks":
                inference_meta.get("num_chunks", 1),
            "Rows per chunk":
                rows_per_chunk_str,
            "Ensemble method":
                inference_meta.get("ensemble_method", "N/A"),
            "Threshold":
                inference_meta.get("prediction_threshold", 0.5),
            "Models used":
                ", ".join(inference_meta.get("models_used", ["tabpfn_binary_model"])),
            "Total inference time (s)":
                round(inf_sec, 2),
            "Average model inference time (s)":
                round(inference_meta.get("avg_model_inference_seconds", 0.0), 2),
            "Throughput (flows/sec)":
                inf_rps,
        }

    if (
        ground_truth_path is not None
        and labeled_flows is not None
    ):

        matched_mask = (
            labeled_flows[
                "label"
            ] != -1
        )

        matched_count = int(
            matched_mask.sum()
        )

        unmatched_count = int(
            len(labeled_flows)
            - matched_count
        )

        report_data[
            "ground_truth"
        ] = {

            "file":
                str(
                    ground_truth_path
                ),

            "matched_flows":
                matched_count,

            "unmatched_flows":
                unmatched_count,

            "match_rate_percentage":
                round(
                    matched_count
                    / len(labeled_flows)
                    * 100,
                    2,
                )
                if len(labeled_flows)
                else 0.0,
        }

    if metrics is not None:

        report_data[
            "evaluation"
        ] = {

            "accuracy":
                metrics.get(
                    "accuracy"
                ),

            "precision":
                metrics.get(
                    "precision"
                ),

            "recall":
                metrics.get(
                    "recall"
                ),

            "f1_score":
                metrics.get(
                    "f1_score"
                ),

            "roc_auc":
                metrics.get(
                    "roc_auc"
                ),

            "true_negatives":
                metrics.get(
                    "true_negatives"
                ),

            "false_positives":
                metrics.get(
                    "false_positives"
                ),

            "false_negatives":
                metrics.get(
                    "false_negatives"
                ),

            "true_positives":
                metrics.get(
                    "true_positives"
                ),
        }

    if predicted_categories is not None or multiclass_predictions is not None:
        import numpy as np
        from tabpfn_nids.evaluation.metrics import decode_attack_category, get_attack_category_mapping

        mapping = get_attack_category_mapping()

        if predicted_categories is not None:
            raw_cats = list(predicted_categories)
        elif multiclass_predictions is not None:
            raw_cats = list(multiclass_predictions)
        else:
            raw_cats = []

        # Guaranteed decode to readable attack category names
        decoded_cats = [decode_attack_category(c) for c in raw_cats]

        if class_names:
            base_classes = [decode_attack_category(c) for c in class_names]
        else:
            base_classes = list(mapping.values())

        # Deduplicate while preserving order
        seen_cls = set()
        all_classes = []
        for c in base_classes:
            if c not in seen_cls:
                seen_cls.add(c)
                all_classes.append(c)

        category_counts = {}
        for c in all_classes:
            category_counts[c] = sum(1 for x in decoded_cats if x == c)
        for x in decoded_cats:
            if x not in category_counts:
                category_counts[x] = sum(1 for y in decoded_cats if y == x)

        category_percentages = {
            cat: round((cnt / total_predictions) * 100, 2) if total_predictions else 0.0
            for cat, cnt in category_counts.items()
        }

        # Malicious attack categories (strictly non-Normal with flows > 0)
        detected_attacks = [
            cat for cat, cnt in category_counts.items()
            if str(cat).lower() != "normal" and cnt > 0
        ]
        total_malicious_flows = sum(
            cnt for cat, cnt in category_counts.items()
            if str(cat).lower() != "normal"
        )
        total_normal_flows = category_counts.get("Normal", 0)

        if prediction_statuses is not None:
            statuses = [str(s).lower() for s in prediction_statuses]
            classified_cnt = sum(1 for s in statuses if s == "classified")
            uncertain_cnt = sum(1 for s in statuses if s == "uncertain")
            unknown_cnt = sum(1 for s in statuses if s == "unknown")
        else:
            classified_cnt = total_predictions
            uncertain_cnt = 0
            unknown_cnt = 0

        avg_conf = None
        if category_confidences is not None and len(category_confidences) > 0:
            avg_conf = round(float(np.mean(category_confidences)), 4)

        report_data["multiclass"] = {
            "task": "Multiclass Attack Identification",
            "classes": all_classes,
            "num_classes": len(all_classes),
            "attack_categories_detected_count": len(detected_attacks),
            "attack_categories_detected": detected_attacks,
            "categories_detected": detected_attacks,
            "category_counts": category_counts,
            "category_percentages": category_percentages,
            "malicious_flows_count": total_malicious_flows,
            "normal_flows_count": total_normal_flows,
            "total_flows_count": total_predictions,
            "status_counts": {
                "classified": classified_cnt,
                "uncertain": uncertain_cnt,
                "unknown": unknown_cnt,
            },
            "classified_count": classified_cnt,
            "uncertain_count": uncertain_cnt,
            "unknown_count": unknown_cnt,
            "uncertain_rate": round(uncertain_cnt / total_predictions, 4) if total_predictions else 0.0,
            "uncertain_rate_percentage": round(uncertain_cnt / total_predictions * 100, 2) if total_predictions else 0.0,
            "uncertainty_threshold": uncertainty_threshold,
            "avg_confidence": avg_conf,
            "supported_classes": all_classes,
            "unsupported_classes": ["Analysis"],
            "unsupported_class_reasons": {
                "Analysis": "Analysis attack category had 0 flow instances in the extracted UNSW-NB15 PCAP slice (timestamps 1421927376 to 1421928304); the earliest Analysis attacks in UNSW-NB15 occur at timestamp 1421930602, so it could not be trained."
            },
            "disclaimer": "Model classifications are statistical machine-learning estimates based on learned feature patterns, not verified forensic proof.",
        }
        report_data["multiclass_summary"] = report_data["multiclass"]

    if multiclass_metrics is not None:
        report_data["multiclass_evaluation"] = multiclass_metrics

    return report_data

def save_docx_report(
    report_data: dict,
    output_path: Path,
) -> Path:
    """Save the PCAP analysis report as a Word document."""

    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.shared import Pt

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    document = Document()

    # -----------------------------------------------------
    # Title
    # -----------------------------------------------------

    title = document.add_heading(
        "PCAP Analysis Report",
        level=0,
    )

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    pcap = report_data.get(
        "pcap",
        {},
    )

    extraction = report_data.get(
        "extraction",
        {},
    )

    features = report_data.get(
        "features",
        {},
    )

    model = report_data.get(
        "model",
        {},
    )

    predictions = report_data.get(
        "predictions",
        {},
    )

    performance = report_data.get(
        "performance",
        {},
    )

    ground_truth = report_data.get(
        "ground_truth"
    )

    evaluation = report_data.get(
        "evaluation"
    )

    inference = report_data.get(
        "inference"
    )

    multiclass = report_data.get(
        "multiclass"
    )

    multiclass_evaluation = report_data.get(
        "multiclass_evaluation"
    )

    generated_at = report_data.get(
        "generated_at",
        "N/A",
    )

    document.add_paragraph(
        f"Generated: {generated_at}"
    )

    # -----------------------------------------------------
    # Helper
    # -----------------------------------------------------

    def add_table(
        title,
        data,
    ):

        document.add_heading(
            title,
            level=1,
        )

        table = document.add_table(
            rows=1,
            cols=2,
        )

        table.alignment = (
            WD_TABLE_ALIGNMENT.CENTER
        )

        table.style = (
            "Table Grid"
        )

        header = table.rows[0].cells

        header[0].text = "Property"
        header[1].text = "Value"

        for key, value in data.items():

            row = table.add_row().cells

            row[0].text = str(key)

            if isinstance(
                value,
                float,
            ):
                value = f"{value:.4f}"

            row[1].text = str(value)

    # -----------------------------------------------------
    # 1. PCAP Information
    # -----------------------------------------------------

    add_table(
        "1. PCAP Information",
        pcap,
    )

    # -----------------------------------------------------
    # 2. Packet and Flow Extraction
    # -----------------------------------------------------

    add_table(
        "2. Packet and Flow Extraction",
        extraction,
    )

    # -----------------------------------------------------
    # 3. Feature Engineering
    # -----------------------------------------------------

    feature_summary = {
        "Flow rows":
            features.get(
                "flow_rows",
                "N/A",
            ),

        "ML feature count":
            features.get(
                "feature_count",
                "N/A",
            ),

        "Total columns after engineering":
            features.get(
                "total_columns_after_engineering",
                "N/A",
            ),
    }

    add_table(
        "3. Feature Engineering",
        feature_summary,
    )

    # -----------------------------------------------------
    # 4. Feature Schema
    # -----------------------------------------------------

    document.add_heading(
        "4. Model Feature Schema",
        level=1,
    )

    feature_columns = features.get(
        "feature_columns",
        [],
    )

    for index, feature in enumerate(
        feature_columns,
        start=1,
    ):

        document.add_paragraph(
            f"{index}. {feature}"
        )

    # -----------------------------------------------------
    # 5. TabPFN Model
    # -----------------------------------------------------

    add_table(
        "5. TabPFN Model",
        model,
    )

    if inference:
        add_table(
            "Inference Architecture & Execution",
            inference,
        )

    # -----------------------------------------------------
    # 6. Prediction Results
    # -----------------------------------------------------

    add_table(
        "6. Prediction Results",
        predictions,
    )

    if multiclass is not None:
        document.add_heading("Multiclass Attack Category Distribution", level=1)
        cat_table = document.add_table(rows=1, cols=3)
        cat_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cat_table.style = "Table Grid"
        h_cells = cat_table.rows[0].cells
        h_cells[0].text = "Attack Category"
        h_cells[1].text = "Flow Count"
        h_cells[2].text = "Percentage (%)"

        cat_counts = multiclass.get("category_counts", {})
        cat_pcts = multiclass.get("category_percentages", {})
        for cat, cnt in cat_counts.items():
            r_cells = cat_table.add_row().cells
            r_cells[0].text = str(cat)
            r_cells[1].text = str(cnt)
            pct_val = cat_pcts.get(cat, 0.0)
            r_cells[2].text = f"{pct_val:.2f}%"

        status_data = {
            "Classified Flows": multiclass.get("status_counts", {}).get("classified", "N/A"),
            "Uncertain Flows": multiclass.get("status_counts", {}).get("uncertain", "N/A"),
            "Unknown / Unclassifiable": multiclass.get("status_counts", {}).get("unknown", "N/A"),
            "Uncertainty Rate": f"{multiclass.get('uncertain_rate_percentage', 0.0):.2f}%",
            "Confidence Threshold": f"{multiclass.get('uncertainty_threshold', 0.5):.2f}",
            "Average Confidence": f"{multiclass.get('avg_confidence', 0.0) * 100:.2f}%" if multiclass.get("avg_confidence") is not None else "N/A",
        }
        add_table("Classification Confidence & Reliability", status_data)

        if multiclass.get("unsupported_classes"):
            document.add_paragraph(
                "Benchmark Coverage Note: Supported attack classes (9): "
                + ", ".join(multiclass.get("supported_classes", []))
                + ". Unsupported classes: "
                + ", ".join(multiclass.get("unsupported_classes", []))
                + " (0 flow instances in training dataset slice)."
            )

    if multiclass_evaluation is not None:
        mc_summary = {
            "Accuracy": multiclass_evaluation.get("accuracy", "N/A"),
            "Macro Precision": multiclass_evaluation.get("macro_precision", "N/A"),
            "Macro Recall": multiclass_evaluation.get("macro_recall", "N/A"),
            "Macro F1-Score": multiclass_evaluation.get("macro_f1", "N/A"),
            "Weighted F1-Score": multiclass_evaluation.get("weighted_f1", "N/A"),
            "OvR ROC-AUC": multiclass_evaluation.get("roc_auc_ovr", "N/A"),
        }
        add_table("Multiclass Evaluation Metrics", mc_summary)

        per_class = multiclass_evaluation.get("per_class", {})
        if per_class:
            document.add_heading("Per-Class Classification Metrics", level=2)
            pc_table = document.add_table(rows=1, cols=5)
            pc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            pc_table.style = "Table Grid"
            pc_h = pc_table.rows[0].cells
            pc_h[0].text = "Category"
            pc_h[1].text = "Precision"
            pc_h[2].text = "Recall"
            pc_h[3].text = "F1-Score"
            pc_h[4].text = "Support"
            for c_name, c_m in per_class.items():
                r = pc_table.add_row().cells
                r[0].text = str(c_name)
                r[1].text = f"{c_m.get('precision', 0.0):.4f}"
                r[2].text = f"{c_m.get('recall', 0.0):.4f}"
                r[3].text = f"{c_m.get('f1', 0.0):.4f}"
                r[4].text = str(c_m.get("support", 0))

    # -----------------------------------------------------
    # 7. Performance
    # -----------------------------------------------------

    add_table(
        "7. Processing Performance",
        performance,
    )

    # -----------------------------------------------------
    # 8. Ground Truth
    # -----------------------------------------------------

    if ground_truth is not None:

        add_table(
            "8. Ground Truth Matching",
            ground_truth,
        )

    # -----------------------------------------------------
    # 9. Evaluation
    # -----------------------------------------------------

    if evaluation is not None:

        add_table(
            "9. Evaluation Metrics",
            evaluation,
        )

    # -----------------------------------------------------
    # 10. Final Summary
    # -----------------------------------------------------

    document.add_heading(
        "10. Analysis Summary",
        level=1,
    )

    total = predictions.get(
        "total",
        0,
    )

    normal = predictions.get(
        "normal_count",
        0,
    )

    attack = predictions.get(
        "attack_count",
        0,
    )

    document.add_paragraph(
        f"The analyzed PCAP contained "
        f"{total} network flows."
    )

    document.add_paragraph(
        f"The TabPFN-NIDS model classified "
        f"{normal} flows as Normal and "
        f"{attack} flows as Attack."
    )

    if evaluation is not None:

        document.add_paragraph(
            "Ground-truth evaluation was "
            "performed for the flows that "
            "successfully matched the supplied "
            "ground-truth dataset."
        )

    else:

        document.add_paragraph(
            "No ground-truth dataset was supplied "
            "for this analysis, so prediction "
            "metrics against ground truth were "
            "not calculated."
        )

    # -----------------------------------------------------
    # Footer
    # -----------------------------------------------------

    section = document.sections[0]

    footer = section.footer.paragraphs[0]

    footer.text = (
        "Generated by the TabPFN-NIDS "
        "PCAP Analysis Pipeline"
    )

    footer.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    # -----------------------------------------------------
    # Font
    # -----------------------------------------------------

    for paragraph in document.paragraphs:

        for run in paragraph.runs:

            run.font.name = "Arial"
            run.font.size = Pt(10)

    # -----------------------------------------------------
    # Save
    # -----------------------------------------------------

    document.save(
        output_path
    )

    return output_path